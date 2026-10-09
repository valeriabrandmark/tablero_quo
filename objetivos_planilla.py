"""La planilla de objetivos de los vendedores: leerla, validarla y armarla.

Es PURO: ni red ni base. La carga a Postgres esta en objetivos.py, que importa
esto. Va aparte por lo mismo que crm.py o relleno.py: los scripts que escriben
abren la base al importarse, y las reglas de la planilla tienen que poder
probarse en el pull request sin credenciales (probar_objetivos.py).

============================================================================
 COMO ES LA PLANILLA
============================================================================

Un archivo por mes comercial en objetivos_mensuales/, igual que los costos:

    objetivos_mensuales/2026-10 objetivos.xlsx

El " objetivos" del nombre es para no confundirlo con la lista de costos, que
se llama 2026-10.xlsx a secas: los dos andan juntos por el escritorio y las
descargas antes de subirse, y el nombre es lo unico que los distingue de
lejos. Ver mes_del_archivo.

Hoja "Objetivos", una fila por objetivo:

    MES_COMERCIAL  2026-10. Tiene que coincidir con el nombre del archivo.
    NOMBRE         como se ve en el tablero: "VASELINE LIP 4.8 G"
    TIPO           SKU      un SKU solo
                   MIX      varios SKUs que se miden SUMADOS
                   MARCA    una marca entera (AVENO)
                   EMPRESA  todas las ventas de una empresa (BRANDMARK, NOA)
    SKU            el SKU, o los SKUs del MIX separados por coma, o la marca,
                   o la empresa
    MEDIDA         UNIDADES (si se deja vacia), FACTURACION o CLIENTES
    VENDEDORES     a quienes aplica: "SILVIO, GERMAN", o TODOS
    CANTIDAD       el objetivo de CADA vendedor de la fila
    NOTA           opcional, para acordarse de por que

Si dos vendedores tienen numeros distintos para el mismo objetivo, van en dos
filas con el mismo NOMBRE. Es lo que paso en 2026-08: GERMAN tenia la mitad
que los demas en todo.

============================================================================
 EL NOMBRE ES EL OBJETIVO, Y NO SE LE PUEDE CAMBIAR EL CONTENIDO
============================================================================

En la base el grupo (gold.objetivos_grupo) no tiene mes: "X5 INSECTICIDAS" es
el mismo grupo en agosto y en octubre. Y el avance se calcula en vivo contra
gold.fact_ventas. Entonces si en octubre "X5 INSECTICIDAS" pasa a sumar un SKU
mas, el avance de AGOSTO tambien cambia, sin que nadie toque agosto.

Ya paso de cerca: en 2026-09 el mix de insecticidas de agosto se partio en
"X5 INSECTICIDA MATA MOSCAS 360 CC" y "X5 DESINFECTANTE DOY PACK 500 CC". Se
hizo bien, con nombres nuevos. Esta planilla lo hace obligatorio: el mismo
NOMBRE con otros SKUs, otro TIPO u otra MEDIDA es un error, y el mensaje pide
otro nombre.
"""

import glob
import hashlib
import os
import re
import unicodedata
from datetime import date, datetime

CARPETA = "objetivos_mensuales"
HOJA = "Objetivos"

# Sube cada vez que cambia lo que objetivos.py escribe. Entra en la huella para
# que --si-cambio recargue una vez aunque ningun Excel se haya tocado: es lo
# mismo que VERSION_ESQUEMA de costos.py.
VERSION_ESQUEMA = 1

COLUMNAS = ["MES_COMERCIAL", "NOMBRE", "TIPO", "SKU", "MEDIDA",
            "VENDEDORES", "CANTIDAD", "NOTA"]
OBLIGATORIAS = ["MES_COMERCIAL", "NOMBRE", "TIPO", "SKU", "VENDEDORES", "CANTIDAD"]

# TIPO de la planilla -> criterio de gold.objetivos_grupo. SKU y MIX son el
# mismo criterio en la base; la planilla los separa porque es la pregunta que
# se hace la comercial ("es un mix o no"), y porque un SKU con tres codigos
# casi siempre es un MIX al que le falto la palabra.
CRITERIO_DE_TIPO = {"SKU": "sku", "MIX": "sku", "MARCA": "marca", "EMPRESA": "empresa"}

MEDIDAS = {"UNIDADES": "unidades", "FACTURACION": "facturacion", "CLIENTES": "clientes"}

# Los vendedores que tienen pagina de objetivos. ES COPIA de
# VENDEDORES_OBJETIVOS en mi-tablero-app/lib/constantes.ts: un objetivo para
# alguien que no esta ahi se cargaria bien y no se veria en ningun lado. Si se
# suma un vendedor, se suma en los dos lados.
VENDEDORES_CON_PAGINA = ["SILVIO", "GERMAN", "PABLO", "RICARDO"]
TODOS = "TODOS"

# Los nombres exactos de las empresas, como los escribe modelo.py (EMPRESAS) en
# gold.fact_ventas. El objetivo matchea por igualdad, asi que "Quo Marketing"
# sin el SRL daria avance cero sin ningun error.
EMPRESAS = ["Quo Marketing SRL", "Noa Comercial SRL", "Presupuesto QUO", "Presupuesto Noa"]

# Los atajos. Los presupuestos SI cuentan para el objetivo, igual que en el
# tablero de Data Studio del que salio la pagina: BRANDMARK son las dos de Quo.
ATAJOS_EMPRESA = {
    "BRANDMARK": ["Quo Marketing SRL", "Presupuesto QUO"],
    "QUO": ["Quo Marketing SRL", "Presupuesto QUO"],
    "NOA": ["Noa Comercial SRL", "Presupuesto Noa"],
}

# El piso del tablero (FECHA_CORTE de modelo.py es el 06/05/2026): antes de eso
# no hay ventas en gold.fact_ventas y el objetivo daria cero siempre.
MES_MINIMO = "2026-05"

PATRON_MES = re.compile(r"\d{4}-\d{2}")
# "2026-10 objetivos", y tambien "2026-10_objetivos", "2026-10-Objetivos" o
# "2026-10" pelado. Lo que NO entra es "2026-10-18": eso es una lista de costos
# de media de mes que se subio a la carpeta equivocada, y tomarla por los
# objetivos de octubre seria peor que ignorarla.
PATRON_ARCHIVO = re.compile(r"(\d{4}-\d{2})(?:[ _-]+objetivos)?", re.IGNORECASE)
SEPARADORES = re.compile(r"\s*[,;/\n]\s*")
# 45.000.000 -> miles con punto, a la argentina.
MILES_CON_PUNTO = re.compile(r"\d{1,3}(\.\d{3})+")


def normalizar(texto):
    """Mayusculas, sin tildes y con un espacio entre palabras.

    "Germán" y "GERMAN" tienen que ser el mismo vendedor: la base guarda
    GERMAN, y una tilde en la planilla lo dejaria sin objetivo.
    """
    sin_tildes = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return " ".join(sin_tildes.upper().split())


def limpio(texto):
    """El texto tal cual, pero sin espacios de mas. Para el NOMBRE, que se ve."""
    return " ".join(str(texto).split())


def vacio(valor):
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


def leer_mes(valor):
    """El mes comercial como 'AAAA-MM', o None si no se entiende.

    Excel convierte "2026-10" en una FECHA apenas se la escribe, aunque la
    plantilla tenga la columna en formato texto si se pega de otro lado. Esa
    fecha llega como datetime y es el mismo mes, asi que se acepta.
    """
    if isinstance(valor, (datetime, date)):
        return f"{valor.year:04d}-{valor.month:02d}"
    if vacio(valor):
        return None
    s = str(valor).strip()
    return s if PATRON_MES.fullmatch(s) else None


def leer_cantidad(valor):
    """El objetivo como numero, o None si no es un numero.

    Llega de tres formas: numero de Excel (45000000), texto a la argentina
    ("45.000.000" o "1.234,5") o texto con punto decimal ("12.5").

    El caso que obliga a no reusar `limpiar_numero` de costos.py es el
    primero de texto: "45.000.000" ahi no tiene coma, se lo pasa a float()
    tal cual y da None. Y un objetivo de facturacion se escribe asi.
    """
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    if vacio(valor):
        return None
    s = str(valor).strip().replace("$", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif MILES_CON_PUNTO.fullmatch(s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def leer_vendedores(valor):
    """('SILVIO', 'GERMAN'), o un error. TODOS son los cuatro con pagina."""
    if vacio(valor):
        return None, "VENDEDORES esta vacia"
    nombres = [normalizar(v) for v in SEPARADORES.split(str(valor)) if v.strip()]
    if TODOS in nombres:
        if len(nombres) > 1:
            return None, "VENDEDORES dice TODOS y ademas nombres: o uno o lo otro"
        return list(VENDEDORES_CON_PAGINA), None
    desconocidos = [n for n in nombres if n not in VENDEDORES_CON_PAGINA]
    if desconocidos:
        return None, (f"vendedor {', '.join(desconocidos)} no tiene pagina de objetivos "
                      f"(valen {', '.join(VENDEDORES_CON_PAGINA)} o {TODOS})")
    repetidos = sorted({n for n in nombres if nombres.count(n) > 1})
    if repetidos:
        return None, f"{', '.join(repetidos)} esta dos veces en VENDEDORES"
    return nombres, None


def leer_items(tipo, valor):
    """Lo que va en gold.objetivos_grupo_item, o un error.

    Los SKUs van en mayusculas. Las marcas tambien, pero recien objetivos.py
    las pasa a como esta escrita en el catalogo, que es lo que matchea. Las
    empresas salen con su nombre exacto de EMPRESAS.
    """
    if vacio(valor):
        return None, "SKU esta vacia"
    partes = [p for p in SEPARADORES.split(str(valor).strip()) if p.strip()]

    if tipo == "SKU":
        if len(partes) != 1:
            return None, (f"TIPO SKU con {len(partes)} SKUs: si se miden sumados, "
                          "el TIPO es MIX")
        return [normalizar(partes[0])], None

    if tipo == "MIX":
        skus = [normalizar(p) for p in partes]
        if len(set(skus)) < 2:
            return None, "un MIX tiene que tener al menos dos SKUs distintos (si es uno, TIPO SKU)"
        return sorted(set(skus)), None

    if tipo == "MARCA":
        if len(partes) != 1:
            return None, "TIPO MARCA lleva una sola marca: si son varias, una fila por marca"
        return [normalizar(partes[0])], None

    # EMPRESA
    por_nombre = {normalizar(e): e for e in EMPRESAS}
    empresas = []
    for p in partes:
        clave = normalizar(p)
        if clave in ATAJOS_EMPRESA:
            empresas.extend(ATAJOS_EMPRESA[clave])
        elif clave in por_nombre:
            empresas.append(por_nombre[clave])
        else:
            return None, (f"empresa '{p}' no existe (valen BRANDMARK, NOA o "
                          f"{', '.join(EMPRESAS)})")
    return sorted(set(empresas)), None


def leer_hoja(nombre_archivo, filas):
    """Valida UN archivo. `filas` son las filas de la hoja, con encabezado.

    Devuelve (objetivos, grupos, errores):

      objetivos  {(vendedor, grupo): cantidad}
      grupos     {grupo: {"criterio", "metrica", "items", "nota", "fila"}}
      errores    lista de textos con el numero de fila de Excel, para que se
                 pueda ir directo a la celda

    Junta TODOS los errores antes de devolver: cargar la planilla y enterarse
    de a uno por corrida, cada hora y media, es una manana entera.
    """
    errores = []
    objetivos, grupos = {}, {}
    mes_archivo = mes_del_archivo(nombre_archivo)

    filas = list(filas)
    if not filas:
        return {}, {}, [f"{nombre_archivo}: la hoja {HOJA} esta vacia"]

    encabezado = [normalizar(c) if not vacio(c) else "" for c in filas[0]]
    faltan = [c for c in OBLIGATORIAS if c not in encabezado]
    if faltan:
        return {}, {}, [f"{nombre_archivo}: faltan las columnas {', '.join(faltan)} "
                        f"en la fila 1 de la hoja {HOJA}"]
    pos = {c: encabezado.index(c) for c in COLUMNAS if c in encabezado}

    # El NOMBRE sin mayusculas ni tildes, para que "Aveno" y "AVENO" no
    # terminen siendo dos objetivos distintos en el tablero.
    nombre_por_clave = {}

    for n, fila in enumerate(filas[1:], start=2):
        celda = {c: (fila[i] if i < len(fila) else None) for c, i in pos.items()}
        if all(vacio(v) for v in celda.values()):
            continue
        donde = f"{nombre_archivo} fila {n}"

        mes = leer_mes(celda["MES_COMERCIAL"])
        if mes is None:
            errores.append(f"{donde}: MES_COMERCIAL '{celda['MES_COMERCIAL']}' no es AAAA-MM")
        elif mes != mes_archivo:
            # El error de siempre: copiar el archivo del mes pasado, cambiarle
            # el nombre y olvidarse de la columna.
            errores.append(f"{donde}: MES_COMERCIAL dice {mes} pero el archivo es "
                           f"el de {mes_archivo}")

        if vacio(celda["NOMBRE"]):
            errores.append(f"{donde}: NOMBRE esta vacio")
            continue
        grupo = limpio(celda["NOMBRE"])
        clave = normalizar(grupo)
        if clave in nombre_por_clave and nombre_por_clave[clave] != grupo:
            errores.append(f"{donde}: '{grupo}' y '{nombre_por_clave[clave]}' son el mismo "
                           "NOMBRE escrito distinto: tiene que ir igual en todas las filas")
            continue
        nombre_por_clave[clave] = grupo

        tipo = normalizar(celda["TIPO"]) if not vacio(celda["TIPO"]) else ""
        if tipo not in CRITERIO_DE_TIPO:
            errores.append(f"{donde}: TIPO '{celda['TIPO']}' no existe "
                           f"(valen {', '.join(CRITERIO_DE_TIPO)})")
            continue

        medida_texto = celda.get("MEDIDA")
        medida = "UNIDADES" if vacio(medida_texto) else normalizar(medida_texto)
        if medida not in MEDIDAS:
            errores.append(f"{donde}: MEDIDA '{medida_texto}' no existe "
                           f"(valen {', '.join(MEDIDAS)})")
            continue

        items, error = leer_items(tipo, celda["SKU"])
        if error:
            errores.append(f"{donde}: {error}")
            continue

        vendedores, error = leer_vendedores(celda["VENDEDORES"])
        if error:
            errores.append(f"{donde}: {error}")
            continue

        cantidad = leer_cantidad(celda["CANTIDAD"])
        if cantidad is None:
            errores.append(f"{donde}: CANTIDAD '{celda['CANTIDAD']}' no es un numero")
            continue
        if cantidad < 0:
            errores.append(f"{donde}: CANTIDAD negativa")
            continue

        definicion = {"criterio": CRITERIO_DE_TIPO[tipo], "metrica": MEDIDAS[medida],
                      "items": items}
        nota = None if vacio(celda.get("NOTA")) else limpio(celda["NOTA"])
        previo = grupos.get(grupo)
        if previo is None:
            grupos[grupo] = {**definicion, "nota": nota, "fila": n}
        else:
            distinto = diferencia(previo, definicion)
            if distinto:
                errores.append(f"{donde}: '{grupo}' ya aparece en la fila {previo['fila']} "
                               f"con otro {distinto}. Las filas del mismo NOMBRE solo "
                               "pueden cambiar VENDEDORES y CANTIDAD")
                continue
            if nota and not previo["nota"]:
                previo["nota"] = nota

        for v in vendedores:
            if (v, grupo) in objetivos:
                errores.append(f"{donde}: {v} ya tiene objetivo de '{grupo}' en otra fila")
                continue
            objetivos[(v, grupo)] = cantidad

    if not objetivos and not errores:
        # Un archivo sin filas BORRARIA el mes entero de la base. Si de verdad
        # un mes no tiene objetivos, se borra el archivo y listo.
        errores.append(f"{nombre_archivo}: no tiene ningun objetivo. Si el mes no "
                       "lleva objetivos, saca el archivo de la carpeta")
    return objetivos, grupos, errores


def diferencia(a, b):
    """Que tienen distinto dos definiciones de grupo, en palabras de la planilla.
    None si son la misma."""
    if a["criterio"] != b["criterio"]:
        return "TIPO"
    if a["metrica"] != b["metrica"]:
        return "MEDIDA"
    if sorted(a["items"]) != sorted(b["items"]):
        return "SKU"
    return None


def combinar(por_mes):
    """Junta los archivos validados y controla los grupos entre meses.

    `por_mes` es {mes: (objetivos, grupos)}. Devuelve (grupos, errores), con
    los grupos ya listos para gold.objetivos_grupo: con su `orden` y su
    `descripcion`.

    EL ORDEN es el de la planilla: el renglon en el que aparece el grupo en el
    mes MAS NUEVO que lo usa. Asi el tablero muestra los objetivos en el orden
    en que se escribieron, y el mes en curso manda. Antes era un numero que se
    elegia a mano en el insert, y cuatro grupos compartian el 2.

    LA NOTA es la del mes mas nuevo que tenga una.
    """
    errores = []
    final = {}
    # Como en leer_hoja, pero entre archivos: "Aveno" en octubre y "AVENO" en
    # agosto serian dos grupos distintos en la base, y el de octubre arrancaria
    # sin la nota ni el orden del otro.
    escrito = {}
    for mes in sorted(por_mes, reverse=True):
        _, grupos = por_mes[mes]
        orden = {g: i for i, g in enumerate(
            sorted(grupos, key=lambda g: grupos[g]["fila"]), start=1)}
        for grupo, d in grupos.items():
            clave = normalizar(grupo)
            if clave in escrito and escrito[clave][0] != grupo:
                errores.append(
                    f"'{grupo}' en {mes} y '{escrito[clave][0]}' en "
                    f"{escrito[clave][1]} son el mismo NOMBRE escrito distinto: "
                    "tiene que ir igual en todos los meses")
                continue
            escrito.setdefault(clave, (grupo, mes))
            if grupo not in final:
                final[grupo] = {"criterio": d["criterio"], "metrica": d["metrica"],
                                "items": d["items"], "orden": orden[grupo],
                                "descripcion": d["nota"], "mes": mes}
                continue
            ya = final[grupo]
            distinto = diferencia(ya, d)
            if distinto:
                errores.append(
                    f"'{grupo}' tiene otro {distinto} en {mes} que en {ya['mes']}. "
                    "El avance se calcula en vivo, asi que cambiarlo reescribiria el "
                    f"avance de {mes}: el objetivo nuevo tiene que llevar otro NOMBRE")
            elif not ya["descripcion"] and d["nota"]:
                ya["descripcion"] = d["nota"]
    return final, errores


def mes_del_archivo(nombre):
    """'2026-10 objetivos.xlsx' -> '2026-10'. None si el nombre no es de un mes.

    Admite el mes pelado ("2026-10.xlsx"), que es como se llamaban al
    principio: renombrar un archivo no tiene que dejarlo afuera.
    """
    base = os.path.splitext(os.path.basename(nombre))[0].strip()
    m = PATRON_ARCHIVO.fullmatch(base)
    return m.group(1) if m else None


def nombre_de_archivo(mes):
    """Como se llama el archivo de un mes cuando lo arma este codigo."""
    return f"{mes} objetivos.xlsx"


def archivos(carpeta=CARPETA):
    """Los archivos de un mes, y aparte los que no tienen ese nombre.

    PLANTILLA.xlsx vive en la misma carpeta a proposito, para tenerla a mano,
    y se ignora por el nombre. El "~$2026-10 objetivos.xlsx" que deja Excel
    abierto tambien.
    """
    validos, ignorados = [], []
    for ruta in sorted(glob.glob(os.path.join(carpeta, "*.xlsx"))):
        (validos if mes_del_archivo(ruta) else ignorados).append(ruta)
    return validos, ignorados


def meses_repetidos(rutas):
    """Errores por cada mes que tiene mas de un archivo.

    Pasa al renombrar: se sube "2026-10 objetivos.xlsx" y queda tambien el
    "2026-10.xlsx" viejo. Elegir uno de los dos seria adivinar cual es el
    bueno, y si se elige el viejo el tablero muestra objetivos que nadie
    quiso cargar.
    """
    por_mes = {}
    for r in rutas:
        por_mes.setdefault(mes_del_archivo(r), []).append(os.path.basename(r))
    return [f"{mes} tiene dos archivos ({' y '.join(sorted(ns))}): tiene que quedar uno solo"
            for mes, ns in sorted(por_mes.items()) if len(ns) > 1]


def huella(carpeta=CARPETA):
    """Firma del CONTENIDO de los archivos de cada mes, como la de costos.py.

    No la fecha de modificacion: GitHub Actions hace checkout limpio en cada
    corrida y la fecha cambia siempre. Eso hizo que costos.py recargara los
    cinco meses cada dos horas (corrida 418, 11/09/2026).
    """
    h = hashlib.sha256()
    h.update(f"v{VERSION_ESQUEMA}|".encode())
    for ruta in archivos(carpeta)[0]:
        h.update(f"{os.path.basename(ruta)}:{os.path.getsize(ruta)}|".encode())
        with open(ruta, "rb") as f:
            for trozo in iter(lambda: f.read(1024 * 1024), b""):
                h.update(trozo)
    return h.hexdigest()


# ============================================================================
#  ARMAR LA PLANILLA
# ============================================================================
# La plantilla vacia y la exportacion de un mes desde la base salen de la misma
# funcion, asi el archivo que se copia para el mes que viene tiene siempre los
# mismos desplegables y la misma hoja de ayuda.

ANCHOS = {"MES_COMERCIAL": 15, "NOMBRE": 38, "TIPO": 11, "SKU": 34, "MEDIDA": 14,
          "VENDEDORES": 30, "CANTIDAD": 14, "NOTA": 60}

AYUDA = [
    ("COMO SE LLENA", None),
    ("Una fila por objetivo. El archivo se llama con el mes comercial: «2026-10 objetivos.xlsx» "
     "y va en la carpeta objetivos_mensuales del repo tablero_quo.", None),
    ("Al subirlo a main, la corrida siguiente del orquestador (cada hora y media) lo "
     "carga y la pagina de cada vendedor se actualiza sola.", None),
    ("", None),
    ("COLUMNA", "QUE VA"),
    ("MES_COMERCIAL", "El mes comercial, AAAA-MM. Igual al nombre del archivo."),
    ("NOMBRE", "Como se ve en el tablero. Si el mismo objetivo tiene numeros distintos "
               "por vendedor, va en varias filas con el mismo NOMBRE."),
    ("TIPO", "SKU = un SKU solo.  MIX = varios SKUs que se miden sumados.  "
             "MARCA = una marca entera.  EMPRESA = todas las ventas de la empresa."),
    ("SKU", "El SKU (PR01001). En un MIX, todos separados por coma (SS06005, SS06006). "
            "En MARCA, la marca (AVENO). En EMPRESA: BRANDMARK o NOA."),
    ("MEDIDA", "UNIDADES (si se deja vacia), FACTURACION (pesos netos) o CLIENTES "
               "(clientes distintos que compraron)."),
    ("VENDEDORES", "A quienes aplica: SILVIO, GERMAN, PABLO, RICARDO separados por coma, "
                   "o TODOS."),
    ("CANTIDAD", "El objetivo de CADA vendedor de la fila. 45000000 o 45.000.000."),
    ("NOTA", "Opcional. Por que el objetivo es como es."),
    ("", None),
    ("TRES REGLAS", None),
    ("1. Un NOMBRE ya usado no puede cambiar de SKUs, TIPO ni MEDIDA, ni en otro mes: "
     "el avance se calcula en vivo y cambiaria el de los meses viejos. Objetivo "
     "distinto, NOMBRE distinto.", None),
    ("2. El archivo reemplaza el mes entero: lo que no esta en la planilla se borra "
     "del tablero para ese mes. Los otros meses no se tocan.", None),
    ("3. Si hay un error no se carga NADA y la corrida avisa con la fila exacta. "
     "Se puede probar antes con: python objetivos.py --revisar", None),
    ("", None),
    ("PARA ARRANCAR UN MES", None),
    ("Copiar el archivo del mes anterior, renombrarlo («2026-11 objetivos.xlsx»), cambiar "
     "MES_COMERCIAL en todas las filas y ajustar los numeros.", None),
]


def escribir_planilla(ruta, filas):
    """Escribe una planilla de objetivos. `filas` son dicts con las COLUMNAS."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    negrita = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F4E78")
    for i, col in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = negrita, fondo
        ws.column_dimensions[c.column_letter].width = ANCHOS[col]
    for f in filas:
        ws.append([f.get(c) for c in COLUMNAS])
    ws.freeze_panes = "A2"

    # Hasta la fila 500: de sobra para un mes (hoy son 9 o 10 objetivos).
    ultima = 500
    letra = {c: ws.cell(row=1, column=i).column_letter for i, c in enumerate(COLUMNAS, start=1)}

    # MES_COMERCIAL en texto, para que Excel no lo convierta en fecha. Si igual
    # lo convierte (pegado de otro lado), leer_mes lo entiende.
    for r in range(2, ultima + 1):
        ws[f"{letra['MES_COMERCIAL']}{r}"].number_format = "@"
        ws[f"{letra['CANTIDAD']}{r}"].number_format = "#,##0"
        ws[f"{letra['NOTA']}{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    for col, opciones in (("TIPO", list(CRITERIO_DE_TIPO)), ("MEDIDA", list(MEDIDAS))):
        dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle=col, error="Elegi una opcion de la lista")
        dv.add(f"{letra[col]}2:{letra[col]}{ultima}")
        ws.add_data_validation(dv)

    # VENDEDORES no lleva lista cerrada porque admite varios a la vez; el
    # desplegable sugiere los casos de un vendedor y TODOS, y deja escribir.
    dv = DataValidation(type="list",
                        formula1='"' + ",".join(VENDEDORES_CON_PAGINA + [TODOS]) + '"',
                        allow_blank=True, showErrorMessage=False)
    dv.add(f"{letra['VENDEDORES']}2:{letra['VENDEDORES']}{ultima}")
    ws.add_data_validation(dv)

    ayuda = wb.create_sheet("Como llenarla")
    ayuda.column_dimensions["A"].width = 18
    ayuda.column_dimensions["B"].width = 100
    for a, b in AYUDA:
        ayuda.append([a, b])
        r = ayuda.max_row
        if b is None and a:
            ayuda.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ayuda.cell(row=r, column=1).alignment = Alignment(wrap_text=True)
            if a.isupper() or a == "COLUMNA":
                ayuda.cell(row=r, column=1).font = Font(bold=True)
            else:
                ayuda.row_dimensions[r].height = 30
        elif a == "COLUMNA":
            ayuda.cell(row=r, column=1).font = Font(bold=True)
            ayuda.cell(row=r, column=2).font = Font(bold=True)
        else:
            ayuda.cell(row=r, column=1).font = Font(bold=True)
            ayuda.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    wb.save(ruta)


def filas_desde_base(mes, objetivos, grupos):
    """Del contenido de la base a las filas de la planilla de UN mes.

    `objetivos` son tuplas (vendedor, grupo, cantidad); `grupos`, {grupo:
    {"criterio", "metrica", "items", "orden", "descripcion"}}.

    Los vendedores con el mismo numero para el mismo grupo van en una sola
    fila, y si son los cuatro, como TODOS: es como lo escribiria una persona.
    """
    por_cantidad = {}
    for vendedor, grupo, cantidad in objetivos:
        por_cantidad.setdefault((grupo, cantidad), []).append(vendedor)

    def tipo_de(g):
        d = grupos[g]
        if d["criterio"] == "sku":
            return "MIX" if len(d["items"]) > 1 else "SKU"
        return d["criterio"].upper()

    def sku_de(g):
        d = grupos[g]
        items = sorted(d["items"])
        if d["criterio"] == "empresa":
            for atajo in ("BRANDMARK", "NOA"):
                if items == sorted(ATAJOS_EMPRESA[atajo]):
                    return atajo
        return ", ".join(items)

    filas = []
    orden_vendedor = {v: i for i, v in enumerate(VENDEDORES_CON_PAGINA)}
    for (grupo, cantidad), vendedores in sorted(
            por_cantidad.items(),
            key=lambda kv: (grupos[kv[0][0]]["orden"], kv[0][0], -kv[0][1])):
        vendedores = sorted(vendedores, key=lambda v: orden_vendedor.get(v, 99))
        filas.append({
            "MES_COMERCIAL": mes,
            "NOMBRE": grupo,
            "TIPO": tipo_de(grupo),
            "SKU": sku_de(grupo),
            "MEDIDA": grupos[grupo]["metrica"].upper(),
            "VENDEDORES": (TODOS if vendedores == VENDEDORES_CON_PAGINA
                           else ", ".join(vendedores)),
            "CANTIDAD": int(cantidad) if float(cantidad).is_integer() else float(cantidad),
            "NOTA": grupos[grupo].get("descripcion"),
        })
    return filas
