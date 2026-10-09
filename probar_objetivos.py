"""Pruebas de la planilla de objetivos, sin base ni red.

POR QUE EXISTE. Casi todo lo que puede salir mal en un objetivo NO da error en
la base: da avance cero, o le cambia el avance a un mes viejo. Un "Germán" con
tilde, un MIX sin la palabra MIX, el archivo del mes pasado copiado sin cambiar
la columna del mes. Lo que se fija aca es que la planilla los frene antes de
llegar a gold.objetivos, con la fila exacta.

Los casos son los de los objetivos reales de 2026-08 y 2026-09.

    python probar_objetivos.py
"""

import os
import shutil
import sys
import tempfile
from datetime import datetime

from openpyxl import load_workbook

import objetivos_planilla as p

FALLOS = []


def revisar(nombre, obtenido, esperado):
    if obtenido == esperado:
        print(f"OK  {nombre}")
    else:
        print(f"MAL {nombre}\n     esperado: {esperado!r}\n     obtenido: {obtenido!r}")
        FALLOS.append(nombre)


def hay_error(errores, pedazo):
    return any(pedazo in e for e in errores)


ENC = p.COLUMNAS


def hoja(*filas):
    """Filas como las devuelve openpyxl: tuplas, con el encabezado primero."""
    return [tuple(ENC)] + [tuple(f) for f in filas]


# --- Numeros -----------------------------------------------------------------

revisar("cantidad: numero de Excel", p.leer_cantidad(45000000), 45000000.0)
# El que obliga a no usar limpiar_numero de costos.py: sin coma, ahi da None.
revisar("cantidad: miles con punto", p.leer_cantidad("45.000.000"), 45000000.0)
revisar("cantidad: miles con punto, corto", p.leer_cantidad("1.500"), 1500.0)
revisar("cantidad: a la argentina con decimales", p.leer_cantidad("1.234,5"), 1234.5)
revisar("cantidad: punto decimal", p.leer_cantidad("12.5"), 12.5)
revisar("cantidad: con signo pesos", p.leer_cantidad("$ 5.000.000"), 5000000.0)
revisar("cantidad: texto", p.leer_cantidad("mucho"), None)
revisar("cantidad: vacia", p.leer_cantidad("  "), None)
revisar("cantidad: un VERDADERO de Excel no es 1", p.leer_cantidad(True), None)

# --- Mes ---------------------------------------------------------------------

revisar("mes: texto", p.leer_mes("2026-10"), "2026-10")
# Excel convierte "2026-10" en fecha apenas se lo tipea fuera de la plantilla.
revisar("mes: Excel lo hizo fecha", p.leer_mes(datetime(2026, 10, 1)), "2026-10")
revisar("mes: cualquier cosa", p.leer_mes("octubre"), None)

# --- Vendedores --------------------------------------------------------------

revisar("vendedores: con tilde y minuscula", p.leer_vendedores("Germán, silvio"),
        (["GERMAN", "SILVIO"], None))
revisar("vendedores: separados por barra", p.leer_vendedores("PABLO / RICARDO")[0],
        ["PABLO", "RICARDO"])
revisar("vendedores: TODOS son los cuatro con pagina", p.leer_vendedores("todos")[0],
        ["SILVIO", "GERMAN", "PABLO", "RICARDO"])
# IGNACIO existe en SIGMA (004) pero no tiene pagina: se cargaria y no se veria.
revisar("vendedores: uno sin pagina", "no tiene pagina" in p.leer_vendedores("IGNACIO")[1], True)
revisar("vendedores: TODOS y nombres", p.leer_vendedores("TODOS, SILVIO")[0], None)
revisar("vendedores: repetido", p.leer_vendedores("SILVIO, silvio")[0], None)
revisar("vendedores: vacio", p.leer_vendedores(None)[0], None)

# --- SKU / marca / empresa ---------------------------------------------------

revisar("items: SKU suelto", p.leer_items("SKU", " pr01001 "), (["PR01001"], None))
revisar("items: MIX ordenado y sin repetidos",
        p.leer_items("MIX", "SS06008, SS06005; SS06006, SS06005"),
        (["SS06005", "SS06006", "SS06008"], None))
revisar("items: SKU con varios es un MIX al que le falto la palabra",
        "el TIPO es MIX" in p.leer_items("SKU", "PR03031, PR03032")[1], True)
revisar("items: MIX de uno", p.leer_items("MIX", "PR03031")[0], None)
revisar("items: BRANDMARK son las dos de Quo, con presupuesto",
        p.leer_items("EMPRESA", "Brandmark")[0], ["Presupuesto QUO", "Quo Marketing SRL"])
revisar("items: NOA", p.leer_items("EMPRESA", "NOA")[0], ["Noa Comercial SRL", "Presupuesto Noa"])
revisar("items: empresa por su nombre, en minuscula",
        p.leer_items("EMPRESA", "quo marketing srl")[0], ["Quo Marketing SRL"])
# "Quo Marketing" sin SRL no matchea nada en fact_ventas: avance cero sin error.
revisar("items: empresa mal escrita", p.leer_items("EMPRESA", "Quo Marketing")[0], None)
revisar("items: dos marcas en una fila", p.leer_items("MARCA", "AVENO, DOVE")[0], None)

# --- Una hoja entera ---------------------------------------------------------

# Lo de 2026-08: GERMAN con la mitad, en otra fila del mismo NOMBRE.
o, g, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "IMPULSE TRUE LOVE 150 ML", "SKU", "PR01001", None, "SILVIO, PABLO, RICARDO", 480),
    ("2026-10", "IMPULSE TRUE LOVE 150 ML", "SKU", "PR01001", "unidades", "GERMAN", 240, "la mitad"),
    (None, None, None, None, None, None, None, None),
    ("2026-10", "Ventas netas NOA", "EMPRESA", "NOA", "Facturación", "TODOS", "5.000.000"),
))
revisar("hoja: sin errores", e, [])
revisar("hoja: la fila vacia del medio no es un error", len(o), 8)
revisar("hoja: GERMAN con su numero", o[("GERMAN", "IMPULSE TRUE LOVE 150 ML")], 240.0)
revisar("hoja: MEDIDA vacia es UNIDADES", g["IMPULSE TRUE LOVE 150 ML"]["metrica"], "unidades")
revisar("hoja: FACTURACION con tilde", g["Ventas netas NOA"]["metrica"], "facturacion")
revisar("hoja: la nota de la segunda fila", g["IMPULSE TRUE LOVE 150 ML"]["nota"], "la mitad")

# El error de siempre: copiar 2026-09.xlsx a 2026-10.xlsx sin tocar la columna.
_, _, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-09", "AVENO", "MARCA", "AVENO", None, "TODOS", 60)))
revisar("hoja: mes de la columna distinto del archivo",
        hay_error(e, "fila 2: MES_COMERCIAL dice 2026-09 pero el archivo es 2026-10"), True)

# El mismo NOMBRE con otros SKUs en la misma hoja.
_, _, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "VASELINE LIP 4.8 G", "MIX", "PR03031, PR03032", None, "SILVIO", 60),
    ("2026-10", "VASELINE LIP 4.8 G", "SKU", "PR03031", None, "GERMAN", 60)))
revisar("hoja: mismo NOMBRE con otros SKUs",
        hay_error(e, "fila 3: 'VASELINE LIP 4.8 G' ya aparece en la fila 2 con otro SKU"), True)

_, _, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "AVENO", "MARCA", "AVENO", None, "TODOS", 60),
    ("2026-10", "AVENO", "MARCA", "AVENO", None, "GERMAN", 30)))
revisar("hoja: el mismo vendedor dos veces en un objetivo",
        hay_error(e, "fila 3: GERMAN ya tiene objetivo de 'AVENO'"), True)

_, _, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "AVENO", "MARCA", "AVENO", None, "SILVIO", 60),
    ("2026-10", "Aveno", "MARCA", "AVENO", None, "GERMAN", 30)))
revisar("hoja: el mismo NOMBRE con otra mayuscula", hay_error(e, "escrito distinto"), True)

_, _, e = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "AVENO", "MARCA", "AVENO", None, "SILVIO", "sesenta"),
    ("2026-10", "X", "COMBO", "AVENO", None, "SILVIO", 1),
    ("2026-10", "Y", "SKU", "PR01001", "PESOS", "SILVIO", 1)))
revisar("hoja: junta todos los errores, no solo el primero", len(e), 3)

_, _, e = p.leer_hoja("2026-10.xlsx", hoja())
revisar("hoja: sin filas no borra el mes", hay_error(e, "no tiene ningun objetivo"), True)

_, _, e = p.leer_hoja("2026-10.xlsx", [("MES_COMERCIAL", "NOMBRE", "SKU")])
revisar("hoja: faltan columnas", hay_error(e, "faltan las columnas TIPO, VENDEDORES, CANTIDAD"), True)

# --- Entre meses -------------------------------------------------------------

ago = p.leer_hoja("2026-08.xlsx", hoja(
    ("2026-08", "X5 INSECTICIDAS", "MIX", "SS07013, SS07014", None, "TODOS", 200, "los 2 de 360"),
    ))
# Octubre le suma un SKU al mismo NOMBRE: le cambiaria el avance a agosto.
octu = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "X5 INSECTICIDAS", "MIX", "SS07013, SS07014, SS07020", None, "TODOS", 300)))
_, e = p.combinar({"2026-08": ago[:2], "2026-10": octu[:2]})
revisar("meses: mismo NOMBRE con otros SKUs en otro mes",
        hay_error(e, "'X5 INSECTICIDAS' tiene otro SKU en 2026-08.xlsx que en 2026-10.xlsx"), True)

octu = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "Aveno", "MARCA", "AVENO", None, "TODOS", 60)))
ago = p.leer_hoja("2026-08.xlsx", hoja(
    ("2026-08", "AVENO", "MARCA", "AVENO", None, "TODOS", 60)))
_, e = p.combinar({"2026-08": ago[:2], "2026-10": octu[:2]})
revisar("meses: el mismo NOMBRE con otra mayuscula en otro mes",
        hay_error(e, "'AVENO' en 2026-08.xlsx y 'Aveno' en 2026-10.xlsx"), True)

sep = p.leer_hoja("2026-09.xlsx", hoja(
    ("2026-09", "ORAL B CEPILLO MICKEY 2 U", "SKU", "GL05011", None, "SILVIO", 80),
    ("2026-09", "Ventas netas NOA", "EMPRESA", "NOA", "FACTURACION", "TODOS", 5000000)))
octu = p.leer_hoja("2026-10.xlsx", hoja(
    ("2026-10", "Ventas netas NOA", "EMPRESA", "NOA", "FACTURACION", "TODOS", 6000000),
    ("2026-10", "AVENO", "MARCA", "AVENO", None, "TODOS", 60)))
grupos, e = p.combinar({"2026-09": sep[:2], "2026-10": octu[:2]})
revisar("meses: sin errores", e, [])
# El mes mas nuevo manda: en octubre NOA va primero.
revisar("meses: el orden sale del mes mas nuevo", grupos["Ventas netas NOA"]["orden"], 1)
revisar("meses: un grupo de un mes viejo conserva su orden",
        grupos["ORAL B CEPILLO MICKEY 2 U"]["orden"], 1)

# --- Archivos ----------------------------------------------------------------

carpeta = tempfile.mkdtemp(prefix="objetivos-")
try:
    # Ida y vuelta: lo que exporta la base se tiene que leer igual. Es lo que
    # garantiza que la primera carga de 2026-08 y 2026-09 de "sin cambios".
    base_grupos = {
        "IMPULSE TRUE LOVE 150 ML": {"criterio": "sku", "metrica": "unidades",
                                     "items": ["PR01001"], "orden": 1, "descripcion": "SKU suelto"},
        "JARDIN DESODORANTE 25 GR": {"criterio": "sku", "metrica": "unidades",
                                     "items": ["SS06005", "SS06006", "SS06008"], "orden": 2,
                                     "descripcion": None},
        "Ventas netas Brandmark": {"criterio": "empresa", "metrica": "facturacion",
                                   "items": ["Presupuesto QUO", "Quo Marketing SRL"],
                                   "orden": 10, "descripcion": None},
    }
    base = [("SILVIO", "IMPULSE TRUE LOVE 150 ML", 480.0), ("PABLO", "IMPULSE TRUE LOVE 150 ML", 480.0),
            ("RICARDO", "IMPULSE TRUE LOVE 150 ML", 480.0), ("GERMAN", "IMPULSE TRUE LOVE 150 ML", 240.0),
            ("SILVIO", "JARDIN DESODORANTE 25 GR", 480.0)]
    base += [(v, "Ventas netas Brandmark", 45000000.0) for v in p.VENDEDORES_CON_PAGINA]
    ruta = os.path.join(carpeta, "2026-08.xlsx")
    filas = p.filas_desde_base("2026-08", base, base_grupos)
    revisar("exportar: los cuatro con el mismo numero son TODOS",
            [f["VENDEDORES"] for f in filas if f["NOMBRE"] == "Ventas netas Brandmark"], ["TODOS"])
    revisar("exportar: la empresa vuelve como BRANDMARK",
            [f["SKU"] for f in filas if f["NOMBRE"] == "Ventas netas Brandmark"], ["BRANDMARK"])
    p.escribir_planilla(ruta, filas)
    wb = load_workbook(ruta, data_only=True)
    o, g, e = p.leer_hoja("2026-08.xlsx", wb[p.HOJA].iter_rows(values_only=True))
    revisar("ida y vuelta: sin errores", e, [])
    revisar("ida y vuelta: los mismos objetivos", o, {(v, gr): c for v, gr, c in base})
    revisar("ida y vuelta: los mismos grupos",
            {k: p.diferencia(g[k], base_grupos[k]) for k in g}, {k: None for k in base_grupos})
    revisar("plantilla: tiene la hoja de ayuda", "Como llenarla" in wb.sheetnames, True)

    # La huella ignora PLANTILLA.xlsx y cambia con el contenido.
    h1 = p.huella(carpeta)
    p.escribir_planilla(os.path.join(carpeta, "PLANTILLA.xlsx"), [])
    revisar("huella: PLANTILLA no cuenta", p.huella(carpeta), h1)
    revisar("archivos: PLANTILLA va a ignorados",
            [os.path.basename(x) for x in p.archivos(carpeta)[1]], ["PLANTILLA.xlsx"])
    p.escribir_planilla(ruta, filas[:-1])
    revisar("huella: cambia si cambia un archivo", p.huella(carpeta) != h1, True)
finally:
    shutil.rmtree(carpeta)

print()
if FALLOS:
    print(f"{len(FALLOS)} prueba(s) fallaron: {', '.join(FALLOS)}")
    sys.exit(1)
print("Todas las pruebas pasaron.")
