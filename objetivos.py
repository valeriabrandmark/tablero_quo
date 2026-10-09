"""Carga los Excel de objetivos de los vendedores a gold.objetivos.

    objetivos_mensuales/2026-10 objetivos.xlsx  ->  los del mes comercial 2026-10

Como es la planilla y por que, en objetivos_planilla.py. Aca esta lo que
necesita la base.

POR QUE. Hasta octubre de 2026 los objetivos se cargaban con los insert y
update del README de mi-tablero-app, a mano en el editor SQL de Supabase. Eso
dejaba en manos de quien escribia el SQL acordarse de que el vendedor va en
mayusculas exactas, de que el grupo exista antes que el objetivo, y de no
cambiarle los SKUs a un grupo que ya usaba un mes viejo. Ninguna de las tres da
error: dan avance cero, o cambian el avance de agosto sin que nadie toque
agosto. Y el mes 2026-10 arranco el 06/10 sin objetivos cargados.

QUE HACE CADA ARCHIVO. Reemplaza su mes ENTERO en gold.objetivos: lo que no
esta en la planilla se borra de ese mes. Los meses sin archivo no se tocan.
Sacar un archivo de la carpeta NO borra su mes de la base, a proposito: perder
un mes de objetivos por un archivo movido es peor que tener que borrarlo a mano.

LOS GRUPOS NO SE BORRAN NUNCA. gold.objetivos tiene `on delete cascade` contra
gold.objetivos_grupo: borrar un grupo se lleva puestos todos los objetivos que
lo usan, de todos los meses. Un grupo que ya no esta en ninguna planilla queda
en la base sin molestar a nadie.

    python objetivos.py                   carga todo lo que hay en la carpeta
    python objetivos.py --revisar         valida y muestra que cambiaria, sin guardar
    python objetivos.py --si-cambio       lo que corre el orquestador
    python objetivos.py --exportar 2026-09  arma "2026-09 objetivos.xlsx" desde la base
"""

import argparse
import os
import sys

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import text

import estado
import objetivos_planilla as planilla
from conexion import crear_engine

load_dotenv()

CLAVE_ESTADO = "objetivos"


def leer_carpeta(carpeta):
    """Lee y valida todos los archivos de mes. Devuelve (por_mes, errores)."""
    validos, ignorados = planilla.archivos(carpeta)
    for ruta in ignorados:
        print(f"  (se ignora {os.path.basename(ruta)}: el nombre no es "
              "'AAAA-MM objetivos')")

    por_mes, errores = {}, planilla.meses_repetidos(validos)
    for ruta in validos:
        nombre = os.path.basename(ruta)
        mes = planilla.mes_del_archivo(nombre)
        if mes < planilla.MES_MINIMO:
            errores.append(f"{nombre}: el tablero arranca en {planilla.MES_MINIMO}, "
                           "antes no hay ventas contra las que medir")
            continue
        try:
            # data_only: si alguien pone una formula en CANTIDAD, vale lo que
            # muestra Excel y no el texto de la formula.
            wb = load_workbook(ruta, read_only=True, data_only=True)
        except Exception as e:
            errores.append(f"{nombre}: no se pudo abrir ({str(e)[:80]})")
            continue
        if planilla.HOJA not in wb.sheetnames:
            errores.append(f"{nombre}: no tiene la hoja '{planilla.HOJA}'")
            continue
        objetivos, grupos, errs = planilla.leer_hoja(
            nombre, wb[planilla.HOJA].iter_rows(values_only=True))
        wb.close()
        errores.extend(errs)
        por_mes[mes] = (objetivos, grupos)
    return por_mes, errores


def grupos_en_base(con):
    """{grupo: {"criterio", "metrica", "items", "orden", "descripcion"}}"""
    filas = con.execute(text("""
        select g.grupo, g.criterio, g.metrica, g.orden, g.descripcion,
               coalesce(array_agg(i.valor order by i.valor)
                        filter (where i.valor is not null), '{}') as items
          from gold.objetivos_grupo g
          left join gold.objetivos_grupo_item i on i.grupo = g.grupo
         group by 1, 2, 3, 4, 5
    """)).mappings().all()
    return {f["grupo"]: {"criterio": f["criterio"], "metrica": f["metrica"],
                         "items": list(f["items"]), "orden": f["orden"],
                         "descripcion": f["descripcion"]} for f in filas}


def revisar_contra_base(con, grupos, meses):
    """Los controles que necesitan la base. Devuelve una lista de errores.

    Corrige de paso la escritura de las marcas: el objetivo matchea por
    igualdad contra gold.fact_ventas, asi que tiene que ir como la escribe el
    catalogo de SIGMA y no como la tipeo alguien.
    """
    errores = []

    # SKUs: contra el catalogo entero y no contra las ventas, porque un
    # producto nuevo puede tener objetivo antes de su primera venta.
    skus = sorted({i for g in grupos.values() if g["criterio"] == "sku" for i in g["items"]})
    if skus:
        existen = set(con.execute(text(
            "select id::text from bronze.sigma_articulos where id::text = any(:s)"
        ), {"s": skus}).scalars())
        for g, d in grupos.items():
            faltan = [s for s in d["items"] if d["criterio"] == "sku" and s not in existen]
            if faltan:
                errores.append(f"'{g}': el SKU {', '.join(faltan)} no esta en el "
                               "catalogo de SIGMA (bronze.sigma_articulos)")

    # Marcas: la columna plana "marca" viene vacia de la API, el dato esta en
    # "attributes.marca" (ver modelo.py, que de ahi saca la de fact_ventas).
    if any(d["criterio"] == "marca" for d in grupos.values()):
        catalogo = {planilla.normalizar(m): m for m in con.execute(text(
            'select distinct "attributes.marca" from bronze.sigma_articulos '
            'where "attributes.marca" is not null')).scalars()}
        for g, d in grupos.items():
            if d["criterio"] != "marca":
                continue
            marca = catalogo.get(planilla.normalizar(d["items"][0]))
            if marca is None:
                errores.append(f"'{g}': la marca {d['items'][0]} no esta en el catalogo")
            else:
                d["items"] = [marca]

    # Un grupo que la base ya tiene con otra definicion. Si lo usa un mes que
    # NO esta en la carpeta, cargar esto le cambiaria el avance a ese mes sin
    # que nadie lo haya pedido. Los meses que si estan ya se compararon entre
    # si en planilla.combinar.
    actuales = grupos_en_base(con)
    usos = con.execute(text(
        "select grupo, array_agg(distinct mes_comercial order by mes_comercial) "
        "from gold.objetivos where not (mes_comercial = any(:m)) group by grupo"
    ), {"m": list(meses)}).all()
    otros_meses = {g: ms for g, ms in usos}
    for g, d in grupos.items():
        previo = actuales.get(g)
        if previo is None or g not in otros_meses:
            continue
        distinto = planilla.diferencia(previo, d)
        if distinto:
            errores.append(
                f"'{g}' ya existe en la base con otro {distinto} y lo usa "
                f"{', '.join(otros_meses[g])}, que no tiene archivo. Cargarlo cambiaria "
                "el avance de ese mes: el objetivo nuevo tiene que llevar otro NOMBRE")
    return errores


def objetivos_en_base(con, meses):
    filas = con.execute(text(
        "select mes_comercial, vendedor, grupo, cantidad from gold.objetivos "
        "where mes_comercial = any(:m)"), {"m": list(meses)}).all()
    return {(m, v, g): float(c) for m, v, g, c in filas}


def mostrar_cambios(antes, despues):
    """Que cambia en la base, mes por mes. Es lo que hay que mirar en el log."""
    hubo = False
    for mes in sorted({k[0] for k in antes} | {k[0] for k in despues}):
        a = {k[1:]: v for k, v in antes.items() if k[0] == mes}
        d = {k[1:]: v for k, v in despues.items() if k[0] == mes}
        altas = sorted(set(d) - set(a))
        bajas = sorted(set(a) - set(d))
        cambios = sorted(k for k in set(a) & set(d) if a[k] != d[k])
        if not (altas or bajas or cambios):
            print(f"  {mes}: sin cambios ({len(d)} objetivos)")
            continue
        hubo = True
        print(f"  {mes}: {len(altas)} nuevos, {len(cambios)} cambiados, "
              f"{len(bajas)} borrados ({len(d)} objetivos)")
        for v, g in altas:
            print(f"     + {v:8} {g}: {d[(v, g)]:,.0f}")
        for v, g in cambios:
            print(f"     ~ {v:8} {g}: {a[(v, g)]:,.0f} -> {d[(v, g)]:,.0f}")
        for v, g in bajas:
            print(f"     - {v:8} {g} (tenia {a[(v, g)]:,.0f})")
    return hubo


def guardar(con, grupos, por_mes):
    """Todo en la transaccion de `con`: grupos, items y los meses enteros.

    Delete + insert y no upsert para los objetivos, porque una fila que se saco
    de la planilla tiene que irse de la base. Y en la misma transaccion que los
    grupos, para que el tablero nunca vea un mes a medio cargar.
    """
    for g, d in grupos.items():
        con.execute(text("""
            insert into gold.objetivos_grupo (grupo, criterio, metrica, orden, descripcion)
            values (:g, :c, :m, :o, :d)
            on conflict (grupo) do update
               set criterio = excluded.criterio,
                   metrica = excluded.metrica,
                   orden = excluded.orden,
                   descripcion = coalesce(excluded.descripcion, gold.objetivos_grupo.descripcion)
        """), {"g": g, "c": d["criterio"], "m": d["metrica"], "o": d["orden"],
               "d": d["descripcion"]})
        con.execute(text("delete from gold.objetivos_grupo_item where grupo = :g"), {"g": g})
        con.execute(text("insert into gold.objetivos_grupo_item (grupo, valor) "
                         "select :g, unnest(cast(:i as text[]))"),
                    {"g": g, "i": d["items"]})

    for mes, (objetivos, _) in por_mes.items():
        con.execute(text("delete from gold.objetivos where mes_comercial = :m"), {"m": mes})
        con.execute(text(
            "insert into gold.objetivos (mes_comercial, vendedor, grupo, cantidad) "
            "values (:m, :v, :g, :c)"),
            [{"m": mes, "v": v, "g": g, "c": c} for (v, g), c in objetivos.items()])


def cargar(carpeta, solo_revisar=False):
    """Devuelve True si quedo todo bien (o no habia nada), False si hubo errores."""
    por_mes, errores = leer_carpeta(carpeta)
    if not por_mes and not errores:
        print(f"  No hay ningun 'AAAA-MM objetivos.xlsx' en {carpeta}/. No hay nada que hacer.")
        return True

    grupos, errs = planilla.combinar(por_mes)
    errores.extend(errs)

    engine = crear_engine()
    with engine.begin() as con:
        # Aunque ya haya errores: un SKU mal tipeado tiene que salir en la
        # misma lista que la fila mal escrita, no en la corrida siguiente.
        errores.extend(revisar_contra_base(con, grupos, por_mes.keys()))

        if errores:
            # NADA se escribe si hay UN error: cargar los meses que si estan
            # bien y no el que tiene la fila mal dejaria al tablero mostrando
            # objetivos viejos para ese mes, sin que se note.
            print(f"\n  {len(errores)} error(es), no se carga nada:\n")
            for e in errores:
                print(f"   - {e}")
            return False

        despues = {(m, v, g): c for m, (objs, _) in por_mes.items()
                   for (v, g), c in objs.items()}
        print("\n  Lo que cambia en gold.objetivos:")
        mostrar_cambios(objetivos_en_base(con, por_mes.keys()), despues)

        if solo_revisar:
            print("\n  --revisar: no se guardo nada.")
            return True
        guardar(con, grupos, por_mes)

    print(f"\n  Guardado: {len(despues)} objetivos de {', '.join(sorted(por_mes))}, "
          f"{len(grupos)} grupos.")
    return True


def exportar(mes, carpeta):
    """Arma "AAAA-MM objetivos.xlsx" con lo que hay en la base para ese mes.

    Sirve para pasar a planilla lo que se cargo a mano antes de que existiera
    esto: con el archivo exportado en la carpeta, la primera carga da "sin
    cambios" y queda probado que el circuito no pisa nada.
    """
    ruta = os.path.join(carpeta, planilla.nombre_de_archivo(mes))
    ya = [r for r in planilla.archivos(carpeta)[0] if planilla.mes_del_archivo(r) == mes]
    if ya:
        sys.exit(f"Ya existe {ya[0]}: no se pisa. Borralo o movelo primero.")
    with crear_engine().begin() as con:
        grupos = grupos_en_base(con)
        objetivos = con.execute(text(
            "select vendedor, grupo, cantidad from gold.objetivos where mes_comercial = :m"
        ), {"m": mes}).all()
    if not objetivos:
        sys.exit(f"No hay objetivos de {mes} en la base.")
    os.makedirs(carpeta, exist_ok=True)
    planilla.escribir_planilla(ruta, planilla.filas_desde_base(
        mes, [(v, g, float(c)) for v, g, c in objetivos], grupos))
    print(f"Listo: {ruta} ({len(objetivos)} objetivos)")


def main():
    parser = argparse.ArgumentParser(
        description="Carga objetivos_mensuales/'AAAA-MM objetivos.xlsx' a gold.objetivos.")
    parser.add_argument("--revisar", action="store_true",
                        help="Valida y muestra que cambiaria, sin guardar")
    parser.add_argument("--si-cambio", action="store_true",
                        help="No hace nada si ningun .xlsx cambio desde la ultima carga")
    parser.add_argument("--exportar", metavar="AAAA-MM",
                        help="Arma el Excel de ese mes con lo que hay en la base")
    args = parser.parse_args()

    if args.exportar:
        if not planilla.PATRON_MES.fullmatch(args.exportar):
            parser.error(f"'{args.exportar}' no es AAAA-MM")
        exportar(args.exportar, planilla.CARPETA)
        return

    # Mismo esquema que costos.py: el orquestador lo llama en cada corrida y
    # el que decide si hay algo que hacer es la huella.
    if args.si_cambio and not args.revisar:
        huella = planilla.huella()
        try:
            guardada = (estado.leer(CLAVE_ESTADO, {}) or {}).get("huella")
        except Exception as e:
            print(f"  (aviso: no se pudo leer la huella: {str(e)[:60]}) -> recarga")
            guardada = None
        if huella == guardada:
            print("Ningun Excel de objetivos cambio desde la ultima carga. No hay nada que hacer.")
            return
        print("Cambio algun Excel de objetivos: se recarga.")

    if not cargar(planilla.CARPETA, solo_revisar=args.revisar):
        # Sale con error para que el orquestador lo cuente como fallo y avise.
        # La huella NO se guarda: la corrida siguiente lo vuelve a intentar, y
        # sigue avisando hasta que alguien arregle la planilla.
        sys.exit(1)

    # Despues de guardar bien, como en costos.py.
    if args.si_cambio and not args.revisar:
        estado.guardar(CLAVE_ESTADO, {"huella": planilla.huella()})
    print("\n=== LISTO ===")


if __name__ == "__main__":
    main()
